"""Update the original-layout competition report with final public wording.

This script intentionally edits content in place while preserving the original
DOCX structure, styles, tables, images, and section layout.
"""
from __future__ import annotations

from pathlib import Path
from shutil import copy2

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "HealthMonitor_ECG_Arrhythmia_Competition_Report.docx"
BEFORE_BACKUP = (
    ROOT
    / "docs"
    / "report_history"
    / "20260612_before_latest_content_update_original_layout.docx"
)
PROGRESS_DIR = ROOT / "docs" / "progress_history"
BACKUP_ROOT = Path(r"C:\HealthMonitor_PreSubmission_Backup_20260612\current_core_v7")


PARAGRAPH_UPDATES = {
    15: "本作品面向连续心电监护场景，设计并实现了一套基于ECG的心律失常类健康预警系统。系统以多源公开心电数据库为基础，采用患者/记录级隔离的数据划分方式，将过去10分钟单导联ECG序列及RR间期特征转换为轨迹样本，预测未来5分钟内的节律软分布和整体心律失常风险。模型采用ArrhythmiaWarningNet双头结构，结合PTB-XL预训练心电编码器、RR统计特征、当前节律监督、未来趋势软标签和转变窗口加权训练；最终采用三成员logits集成模型，并通过验证集确定EWMA平滑、阈值和连续触发策略后，在锁定测试集上固定应用。测试结果显示，最终模型在锁定测试集上取得future macro-AUROC 0.747、future macro-F1 0.318；未来心律失常事件级报警召回率0.925，incident event recall 0.897，中位提前预警时间290秒，false alert episodes为1.12次/患者小时。AFib辅助方向分数在LTAFDB 83条固定质量记录的冻结外部主分析中取得AUROC 0.721、事件召回0.613和中位提前240秒，84条敏感性分析结论接近，但该方向未接入正式报警。系统还实现了三模型一致性、信号质量、窗口级Grad-CAM和ECG级Integrated Gradients解释，可在演示端展示风险曲线、预警状态和模型关注片段。本作品定位为科研和竞赛展示的早期预警原型，不替代临床诊断；正式能力限定为未来5分钟整体心律失常风险预警，PVC相关风险识别表现较稳定，AFib为外部研究支持的辅助方向，VT/VF和AT/SVT仅作为探索性能力边界呈现。",
    64: "项目数据层使用MIT-BIH Arrhythmia Database、AFDB、SVDB、VFDB等公开心电数据库，并保留PTB-XL作为预训练心电表征来源。PTB-XL在当前系统中作为backbone初始化和心电表征基础，而非直接作为未来预警标签来源。所有正式指标均来自固定数据划分、锁定评估脚本和冻结证据文件，避免在报告中依赖本地绝对路径或临时实验目录。",
    66: "项目已经形成从数据规划、数据审计、模型训练、多模型稳定性评估、验证集选型、锁定测试集评估、报警策略评估、冻结外部评估、API推理服务到三页式Streamlit演示端的完整链路。当前正式策略为overall-risk-v1：以1-P(future Normal)作为整体风险，EWMA alpha为0.65，阈值0.10，连续2个10秒推理点触发正式报警。公开展示端采用competition-demo-v2清单，默认病例为MIT-BIH 119 PVC报警、MIT-BIH 100正常无报警和MIT-BIH 201 AFib辅助方向；223与209仅在能力边界区域说明。",
    80: "推理阶段，FastAPI服务加载正式三成员集成模型，对三个模型logits取平均得到最终概率，并输出原始整体风险、AFib辅助方向分数、信号质量、模型投票数、风险离散度和策略版本。前端与离线评估复用同一冻结策略模块，按模拟时间每10秒更新一次推理；波形浏览与推理频率分离，事件日志仅记录EWMA平滑、阈值0.10和连续2次触发后的正式报警开始或解除。",
    85: "单模型部署简单但随机性较强；单成员最优模型可能受偶然波动影响。当前系统采用三成员logits集成，一方面提升验证集判别能力，另一方面为risk std、类别投票数等不确定性展示提供自然来源。公开报告只保留集成方案和最终指标，不展示内部种子编号、检查点命名或调参轮次。",
    88: "每个模型输入样本包含过去600秒ECG历史，划分为39个30秒窗口，窗口步长15秒，单窗口长度7500个采样点。除原始波形外，系统提取mean RR、SDNN、RMSSD、pNN50、样本熵、Poincare特征和RR变异系数等RR相关特征。标签体系包含Normal、PVC、AFib、VF、VT、AT/SVT六类。训练阶段采用类别权重、转变窗口加权和多候选模型选择策略，转变窗口权重只作用于未来趋势损失，强调“当前与未来不同”的关键样本。",
    91: "核心模型ArrhythmiaWarningNet由窗口级心电编码器、RR特征编码器、序列融合层和两个输出头组成。最终默认报警策略针对future_arrhythmia任务，使用未温度缩放的集成概率，阈值0.10，EWMA平滑alpha=0.65，连续2个窗口触发报警。AFib辅助方向在不改变正式整体预警的前提下，以P(AFib)/[P(AFib)+P(Normal)]构造方向分数；其阈值0.10、EWMA alpha=0.65和k=1均在外部推理前冻结，仅用于证据页和指定病例展示。",
    96: "项目采用患者/记录级划分数据；训练集只用于参数学习；验证集用于选择模型、校准方式、阈值和平滑报警策略；锁定测试集只应用验证集已经确定的模型和策略。该流程减少了数据泄漏和测试集调参造成的乐观偏差。",
    107: "AFib辅助方向首先在验证集上完成候选策略筛选：验证集包含2188个窗口、21个患者/记录单位，其中349个主导未来AFib窗口来自5个记录。未来类别argmax的AFib precision、recall和F1分别为0.251、0.160和0.196；采用P(AFib)/[P(AFib)+P(Normal)]、EWMA alpha=0.65、阈值0.10、k=1后，三项指标分别为0.322、0.653和0.431，AFib事件召回为0.733（11/15），中位提前量280秒，AFib特异误报为3.71次/患者小时。该候选在5次阳性记录留一影响检查中通过4次，但患者级bootstrap区间仍较宽，因此仅进入外部研究验证，不用于锁定测试集选型，也不改变正式整体预警。",
    108: "完整冻结外部评估覆盖Long-Term AF Database全部84条记录。信号审计按清单顺序完成，所有文件大小与头文件一致；记录20的官方头文件初值和校验和与官方信号字节不一致，本地整文件SHA-256与PhysioNet源文件完全相同，因此固定质量主分析纳入其余83条记录，记录20仅进入84条敏感性分析。冻结协议ltafdb-external-v3共生成698256个10秒评估时点，外部阈值未调整。83条主分析中，六分类argmax基线的主导未来AFib recall为0.094、事件recall为0.292、误报episode为2.082次/患者小时；固定AFib方向策略对应为0.409、0.613和2.330，主导AFib AUROC为0.721、AP为0.706，中位提前量240秒。1000次患者级bootstrap的AUROC 95%区间为0.661～0.775，事件recall区间为0.506～0.720。84条敏感性结果与主分析接近，未改变结论。该结果可表述为LTAFDB 83记录主分析、84记录敏感性分析的外部研究支持，但不改变正式整体预警，也不构成临床有效性或全面跨数据库泛化结论。",
    109: "长尾方向审计显示：VT/(VT+Normal)虽达到事件recall 0.667和中位提前215秒，但专项误报2.851次/患者小时，未通过门槛；VF因阳性记录不足停止；VT/VF组合和AT/SVT方向虽在验证筛选中出现一定信号，但在后续固定门控与稀有类别验证中未形成可接入锁定测试和正式报警的稳定证据。因此，提交前不再宣传VT/VF或AT/SVT专项识别能力，相关病例仅作为能力边界展示。",
    110: "系统补充软标签一致评估，以完整未来5分钟类别分布而非仅用argmax硬标签衡量概率输出。验证集2188个窗口的soft cross-entropy为1.336、soft KL为1.010、soft Brier为0.380；未来异常质量的MAE为0.266、RMSE为0.355、平均偏差为0.013。上述指标用于揭示概率分布误差，不替代既有锁定报警指标，也未用于重新选择正式报警阈值。",
    115: "项目已经具备端到端工程链路：后端加载正式三成员集成模型，保持旧接口兼容并明确报告输入状态；前端重构为“监测预警、精选病例、验证证据”三个页面，分别面向使用逻辑、现场演示和评委证据核查。页面统一使用“正式冻结策略”“辅助研究能力”“探索性边界”等审慎术语，并提供信号异常、短时波动、持续报警和伴随症状四类安全提示。AFib方向分数仅在外部研究证据和指定病例中显示，未改变正式整体预警。",
    119: "本作品的核心宣传边界应聚焦于未来5分钟整体心律失常风险预警和PVC相关风险识别。AFib辅助方向已获得LTAFDB 83条固定质量记录的全库外部研究支持，患者级区间和84条敏感性分析均已报告；它可以用于证据页和指定病例的辅助方向展示，但仍不能表述为临床可用能力，也不得替代正式整体报警。独立VT、VF、VT/VF专项方向和AT/SVT均未形成可晋级锁定测试的稳定证据，不得作为已验证类别预警能力宣传。",
    122: "在解释性样例中，MIT-BIH 119记录的PVC目标概率为0.839，所选关键窗口包含5个V/E标注搏动，ECG级归因在标注邻域与背景的平均强度比为5.32，说明解释热区与异常搏动存在较好对应。",
    128: "后续优化应优先围绕外部数据泛化、信号质量控制、前端异常输入提示、真实采集链路和医学安全边界展开。长尾类别方向需要更多阳性记录、稳定事件召回、足够提前量和可接受误报负担同时满足后，才可进入新的专项研究；在此之前不覆盖现有正式数据、模型或锁定结果。系统工程下一步可接入ECG模拟器和串口/USB采集链路，展示信号质量控制与异常输入状态；未经伦理审批的人体采集不作为比赛证据。",
    129: "综上，本作品已经形成包含多数据库训练、患者级评估、事件级报警、软标签一致评估、信号质量、模型一致性、可解释分析、LTAFDB全库冻结外部研究支持和可复现病例回放的心律失常健康预警原型。当前正式能力仍是未来5分钟整体心律失常风险提示；AFib方向为具有全库外部研究支持但未接入正式报警的辅助能力；VT/VF和AT/SVT仍为探索能力边界。系统输出用于风险提示和专业复核，不替代临床诊断。",
}

TABLE_UPDATES = {
    (0, 1, 2): "完成MIT-BIH、AFDB、SVDB、VFDB等数据库的统一读取、重采样、窗口化与数据审计；正式数据无非有限值。",
    (0, 5, 1): "多模型稳定性验证",
    (0, 5, 2): "完成多成员训练与验证，验证集future macro-AUROC约0.731±0.013，future macro-F1约0.278±0.014。",
    (0, 6, 1): "正式集成模型与锁定测试",
    (0, 6, 2): "三成员logits集成在锁定测试集future macro-AUROC 0.747，future macro-F1 0.318。",
    (0, 9, 2): "实现三模型ensemble disagreement，输出risk mean、risk std、模型投票数和风险离散度。",
    (2, 1, 2): "冻结轨迹数据与质量审计清单",
    (2, 4, 2): "验证集选型 + 锁定测试",
    (6, 3, 0): "多成员验证均值",
    (9, 1, 0): "正式冻结策略",
    (9, 2, 0): "患者自适应探索策略",
    (9, 3, 0): "RR gating探索策略",
    (11, 1, 0): "核心代码",
    (11, 1, 1): "healthmonitor/，main.py，dashboard.py",
    (11, 2, 1): "scripts/data/build_dataset_factory.py，datasets/trajectory_official",
    (11, 3, 1): "healthmonitor/dl_model.py，ArrhythmiaWarningNet",
    (11, 4, 1): "scripts/training/train_trajectory.py，scripts/training/多成员训练入口",
    (11, 5, 1): "scripts/evaluation/final_validation_analysis.py，scripts/evaluation/evaluate_alert_policy.py，scripts/external_validation/evaluate_ltafdb_external.py",
    (11, 6, 0): "最终策略",
    (11, 6, 1): "artifacts/policies/final_alert_policy_selected.json，policy_config_version=overall-risk-v1",
    (11, 7, 0): "锁定测试证据",
    (11, 7, 1): "artifacts/evidence/locked_validation_observed_metrics.json，locked_validation_event_lead_time_summary.json",
    (11, 8, 0): "外部研究证据",
    (11, 8, 1): "artifacts/evidence/ltafdb_full_observed_metrics.json，ltafdb_full_patient_bootstrap_ci.csv，ltafdb_quality_sensitivity.csv",
    (11, 9, 0): "API与演示",
    (11, 9, 1): "healthmonitor/main.py，healthmonitor/dashboard.py，artifacts/demo/replay_manifest.json",
    (11, 10, 0): "AFib辅助方向",
    (11, 10, 1): "artifacts/evidence/ltafdb_full_decision.json，artifacts/evidence/ltafdb_full_observed_metrics.json",
    (11, 11, 0): "能力边界与软标签证据",
    (11, 11, 1): "artifacts/evidence/long_tail_directional_decision.json，artifacts/evidence/soft_target_observed_metrics.json，artifacts/demo/mitdb_223.json，artifacts/demo/mitdb_209.json",
}

PROGRESS_SOURCES = [
    "results_v7_round2/ROUND2_SUMMARY.md",
    "results_v7_round3/ROUND3_SUMMARY.md",
    "results_v7_round4/ROUND4_PROGRESS.md",
    "results_v7_round5/ROUND5_OPTIMIZATION_PLAN.md",
    "results_v7_round5/ROUND5_PROGRESS.md",
    "results_v7_round6/ROUND6_PROGRESS.md",
    "results_v7_round7/ROUND7_PLAN.md",
    "results_v7_round7/afib_baseline/AFIB_BASELINE.md",
    "results_v7_round7/ltafdb_external_evaluation/LTAFDB_EXTERNAL_EVALUATION.md",
    "results_v7_round7/ltafdb_pilot_audit/LTAFDB_PILOT_AUDIT.md",
    "results_v7_round7/ltafdb_signal_audit/LTAFDB_SIGNAL_AUDIT.md",
    "results_v7_round8/ROUND8_PROGRESS.md",
    "results_v7_round8/cudb_external_eligibility/CUDB_ELIGIBILITY.md",
    "results_v7_round8/dataset_v8_rare_plan/V8_RARE_SPLIT.md",
    "results_v7_round8/long_tail_directional_audit/DIRECTIONAL_DECISION.md",
    "results_v7_round8/ltafdb_full_external_evaluation/LTAFDB_FULL_EVALUATION.md",
    "results_v7_round8/ltafdb_full_signal_audit/LTAFDB_SIGNAL_AUDIT.md",
    "results_v7_round8/soft_target_evaluation/SOFT_TARGET_EVALUATION.md",
]


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    replace_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        replace_paragraph_text(extra, "")


def update_report() -> None:
    BEFORE_BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BEFORE_BACKUP.exists():
        copy2(REPORT, BEFORE_BACKUP)

    doc = Document(REPORT)
    for idx, text in PARAGRAPH_UPDATES.items():
        replace_paragraph_text(doc.paragraphs[idx], text)

    for (table_idx, row_idx, cell_idx), text in TABLE_UPDATES.items():
        replace_cell_text(doc.tables[table_idx].rows[row_idx].cells[cell_idx], text)

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.subject = ""
    props.keywords = ""
    props.category = ""

    tmp = REPORT.with_name(REPORT.stem + "_content_update_tmp.docx")
    doc.save(tmp)
    tmp.replace(REPORT)


def archive_progress_reports() -> None:
    PROGRESS_DIR.mkdir(parents=True, exist_ok=True)
    entries = []
    for rel in PROGRESS_SOURCES:
        src = BACKUP_ROOT / rel
        if src.exists():
            dest = PROGRESS_DIR / rel.replace("/", "__")
            copy2(src, dest)
            entries.append((rel, dest.name, dest.stat().st_size))
        else:
            entries.append((rel, None, None))

    lines = [
        "# 项目每轮进程汇报归档",
        "",
        "本目录保存项目迭代过程中形成的阶段汇报、专项审计和过程说明，用于回顾历史决策与选择。正式竞赛报告不再展开这些内部轮次叙事，仅保留最终证据和能力边界。",
        "",
        "| 来源路径 | 归档文件 | 大小 |",
        "|---|---:|---:|",
    ]
    for rel, name, size in entries:
        if name is None:
            lines.append(f"| `{rel}` | 未找到 | - |")
        else:
            lines.append(f"| `{rel}` | `{name}` | {size} |")

    index = PROGRESS_DIR / "INDEX.md"
    index.write_text("\n".join(lines) + "\n", encoding="utf-8")

    update_note = PROGRESS_DIR / "20260612_report_content_update_progress.md"
    update_note.write_text(
        "\n".join(
            [
                "# 2026-06-12 报告内容更新进程汇报",
                "",
                "## 本轮目标",
                "- 在根目录正式实验报告原格式、原目录、原排版基础上，仅更新内容。",
                "- 清理正式报告中的内部轮次叙事、本地绝对路径、旧工作区名称、检查点/种子编号等提交前不宜公开的内容。",
                "- 保留每轮修改后产生的进程汇报文件，统一归档到 `docs/progress_history/`，便于后续回顾历史过程与选择。",
                "",
                "## 已执行",
                "- 已在 `docs/report_history/20260612_before_latest_content_update_original_layout.docx` 保存本轮更新前的原版式备份。",
                "- 已将正式报告摘要、研究基础、技术路线、测试结果、能力边界、总结和附录表格改为最终提交口径。",
                "- 已将 AFib 表述统一为 LTAFDB 83记录主分析、84记录敏感性分析的外部研究支持；未接入正式报警。",
                "- 已将 VT/VF、AT/SVT 表述压缩为探索性能力边界，不作为正式验证能力宣传。",
                "- 已把外置备份中的阶段汇报和专项审计 Markdown 复制到 `docs/progress_history/` 并生成索引。",
                "",
                "## 验收要点",
                "- 正式报告不再包含 Round、current_core_v7、C:\\\\、V7、V8、results_v7 等内部路径或轮次词。",
                "- 正式报告仍保持原 DOCX 文件路径：`HealthMonitor_ECG_Arrhythmia_Competition_Report.docx`。",
                "- 历史过程材料不写入正式报告正文，而是作为进程汇报归档保留。",
                "",
            ]
        ),
        encoding="utf-8",
    )


if __name__ == "__main__":
    update_report()
    archive_progress_reports()
    print(REPORT)
    print(PROGRESS_DIR / "INDEX.md")
